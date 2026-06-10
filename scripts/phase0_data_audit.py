from __future__ import annotations

import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

PROJECT = Path(__file__).resolve().parents[1]
REPORT_XLSX = PROJECT / "QuantumQueen_backtest_report.xlsx"
M1_CSV = PROJECT / "xauusd_m1_2016_2026.csv"
M15_CSV = PROJECT / "xauusd_m15_2016_2025.csv"
OUTPUTS = PROJECT / "outputs"


def clean_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, np.generic):
        return value.item()
    return value


def first_row_index(raw: pd.DataFrame, label: str) -> int:
    matches = raw.index[raw[0].astype(str).str.strip().eq(label)]
    if len(matches) == 0:
        raise ValueError(f"Could not find section label: {label}")
    return int(matches[0])


def read_workbook_shape() -> dict:
    wb = load_workbook(REPORT_XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    return {"sheets": wb.sheetnames, "max_row": ws.max_row, "max_column": ws.max_column}


def parse_metadata(raw: pd.DataFrame) -> tuple[dict, dict]:
    metadata = {}
    inputs = {}
    for i in range(0, min(len(raw), 82)):
        row = raw.iloc[i]
        key = clean_value(row.iloc[0])
        value = clean_value(row.iloc[3]) if len(row) > 3 else None
        if isinstance(key, str) and key.strip().endswith(":"):
            metadata[key.strip().rstrip(":")] = value
        if isinstance(value, str) and "=" in value:
            k, v = value.split("=", 1)
            inputs[k.strip()] = v.strip()
    return metadata, inputs


def parse_orders(raw: pd.DataFrame, order_label_idx: int, deals_label_idx: int) -> pd.DataFrame:
    rows = raw.iloc[order_label_idx + 2 : deals_label_idx, :13].copy()
    rows.columns = [
        "open_time",
        "order",
        "symbol",
        "type",
        "volume_raw",
        "blank_1",
        "requested_price",
        "sl",
        "tp",
        "time",
        "blank_2",
        "status",
        "comment",
    ]
    rows = rows[rows["open_time"].notna()].copy()
    rows["open_time"] = pd.to_datetime(rows["open_time"], format="%Y.%m.%d %H:%M:%S", errors="coerce")
    rows["time"] = pd.to_datetime(rows["time"], format="%Y.%m.%d %H:%M:%S", errors="coerce")
    rows["order"] = pd.to_numeric(rows["order"], errors="coerce").astype("Int64")
    rows["requested_price"] = pd.to_numeric(rows["requested_price"], errors="coerce")
    return rows


def parse_deals(raw: pd.DataFrame, deals_label_idx: int) -> pd.DataFrame:
    rows = raw.iloc[deals_label_idx + 2 :, :13].copy()
    rows.columns = [
        "time",
        "deal",
        "symbol",
        "type",
        "entry",
        "volume",
        "price",
        "order",
        "commission",
        "swap",
        "profit",
        "balance",
        "comment",
    ]
    rows = rows[rows["time"].notna()].copy()
    rows = rows[rows["symbol"].eq("XAUUSD")].copy()
    rows["time"] = pd.to_datetime(rows["time"], format="%Y.%m.%d %H:%M:%S", errors="coerce")
    for col in ["deal", "order"]:
        rows[col] = pd.to_numeric(rows[col], errors="coerce").astype("Int64")
    for col in ["volume", "price", "commission", "swap", "profit", "balance"]:
        rows[col] = pd.to_numeric(rows[col], errors="coerce")
    pattern = r"QQ\[(?P<tag_symbol>[^\]]+)\](?P<magic>\d+)\[T(?P<T>\d+)/S(?P<S>\d+)\]"
    extracted = rows["comment"].astype(str).str.extract(pattern)
    rows = pd.concat([rows, extracted], axis=1)
    rows["strategy"] = None
    tagged = rows["T"].notna() & rows["S"].notna()
    rows.loc[tagged, "strategy"] = (
        "T"
        + rows.loc[tagged, "T"].astype(int).astype(str)
        + "/S"
        + rows.loc[tagged, "S"].astype(int).astype(str).str.zfill(2)
    )
    return rows


def csv_summary(path: Path) -> dict:
    sample = pd.read_csv(path, nrows=5)
    times = pd.read_csv(path, usecols=["time"])
    parsed = pd.to_datetime(times["time"], utc=True, errors="coerce")
    return {
        "file": path.name,
        "size_bytes": path.stat().st_size,
        "columns": list(sample.columns),
        "sample_rows": sample.astype(str).to_dict(orient="records"),
        "row_count": int(len(times)),
        "time_min": str(parsed.min()),
        "time_max": str(parsed.max()),
        "duplicate_timestamps": int(times["time"].duplicated().sum()),
        "missing_timestamps": int(parsed.isna().sum()),
    }


def alignment_scores(deals: pd.DataFrame) -> dict:
    m1 = pd.read_csv(M1_CSV, usecols=["time", "high", "low", "spread"])
    m1["time"] = pd.to_datetime(m1["time"], utc=True).dt.tz_localize(None)
    m1 = m1.set_index("time").sort_index()
    trade_deals = deals.dropna(subset=["time", "price"]).copy()
    scores = []
    for tolerance in [0, 0.05, 0.1, 0.25, 0.5]:
        for offset in range(-6, 7):
            aligned_time = trade_deals["time"].dt.floor("min") + pd.Timedelta(hours=offset)
            probe = pd.DataFrame({"time": aligned_time.values, "price": trade_deals["price"].values})
            joined = probe.join(m1[["high", "low", "spread"]], on="time")
            available = joined["high"].notna()
            within = available & (joined["price"] >= joined["low"] - tolerance) & (joined["price"] <= joined["high"] + tolerance)
            distance = np.where(~available, np.nan, np.maximum(joined["low"] - joined["price"], joined["price"] - joined["high"]))
            distance = np.where(distance < 0, 0, distance)
            scores.append(
                {
                    "offset_hours": offset,
                    "tolerance": tolerance,
                    "available": int(available.sum()),
                    "match": int(within.sum()),
                    "match_rate": float(within.sum() / available.sum()) if available.sum() else 0.0,
                    "median_distance": float(np.nanmedian(distance)) if available.sum() else None,
                    "p90_distance": float(np.nanpercentile(distance, 90)) if available.sum() else None,
                }
            )
    entries = trade_deals[trade_deals["entry"].eq("in")].copy()
    entry_scores = []
    for offset in range(-6, 7):
        aligned_time = entries["time"].dt.floor("min") + pd.Timedelta(hours=offset)
        probe = pd.DataFrame({"time": aligned_time.values, "price": entries["price"].values})
        joined = probe.join(m1[["high", "low"]], on="time")
        available = joined["high"].notna()
        within = available & (joined["price"] >= joined["low"] - 0.25) & (joined["price"] <= joined["high"] + 0.25)
        entry_scores.append(
            {
                "offset_hours": offset,
                "tolerance": 0.25,
                "available": int(available.sum()),
                "match": int(within.sum()),
                "match_rate": float(within.sum() / available.sum()) if available.sum() else 0.0,
            }
        )
    best = max([s for s in scores if s["tolerance"] == 0.25], key=lambda s: s["match_rate"])
    best_entry = max(entry_scores, key=lambda s: s["match_rate"])
    return {"all_deal_scores": scores, "entry_scores_tolerance_0_25": entry_scores, "best_tolerance_0_25": best, "best_entry_tolerance_0_25": best_entry}


def value_counts_dict(series: pd.Series) -> dict:
    return {str(k): int(v) for k, v in series.value_counts(dropna=False).sort_index().items()}


def build_report(summary: dict) -> str:
    strategy_rows = "\n".join([f"| {k} | {v} |" for k, v in summary["deals"]["entry_strategy_counts"].items()])
    s_rows = "\n".join([f"| S{int(k):02d} | {v} |" for k, v in summary["deals"]["entry_s_counts"].items()])
    alignment = summary["time_alignment"]
    top_offsets = sorted([s for s in alignment["all_deal_scores"] if s["tolerance"] == 0.25], key=lambda s: s["match_rate"], reverse=True)[:7]
    offset_rows = "\n".join(
        [
            f"| {r['offset_hours']} | {r['available']} | {r['match']} | {r['match_rate']:.4f} | {r['median_distance']:.4f} | {r['p90_distance']:.4f} |"
            for r in top_offsets
        ]
    )
    lines = [
        "# Phase 0 Data Audit Report",
        "",
        "## Executive Summary",
        "",
        f"- Workbook sheets: `{summary['workbook']['sheets']}`.",
        f"- Main sheet dimensions: `{summary['workbook']['max_row']}` rows x `{summary['workbook']['max_column']}` columns.",
        f"- Expert: `{summary['metadata'].get('专家')}`.",
        f"- Symbol: `{summary['metadata'].get('交易品种')}`.",
        f"- Test period field: `{summary['metadata'].get('期间')}`.",
        f"- Magic input: `{summary['inputs'].get('InpMagicNumber')}`.",
        f"- Set input: `{summary['inputs'].get('InpSets')}`.",
        f"- Orders parsed: `{summary['orders']['row_count']}`.",
        f"- Deals parsed: `{summary['deals']['row_count']}`.",
        f"- Entry deals: `{summary['deals']['entry_count']}`; exit deals: `{summary['deals']['exit_count']}`.",
        f"- Every entry deal has an explicit strategy tag: `{summary['deals']['entries_with_strategy']}/{summary['deals']['entry_count']}`.",
        f"- Active strategy tags observed: `{len(summary['deals']['entry_strategy_counts'])}`.",
        "- This report contains explicit `T*/S*` strategy identifiers in entry comments, so Phase 1+ should preserve those labels before attempting behavioral clustering.",
        "",
        "## Timezone / Offset Check",
        "",
        "The user recalled that exported trade-record time is GMT+3. Against the provided M1 CSV, the best alignment is not `-3h`; it is `0h` after flooring trade timestamps to the minute.",
        "",
        f"- Best all-deal offset at tolerance 0.25: `{alignment['best_tolerance_0_25']['offset_hours']}h`, match rate `{alignment['best_tolerance_0_25']['match_rate']:.4f}`.",
        f"- Best entry-only offset at tolerance 0.25: `{alignment['best_entry_tolerance_0_25']['offset_hours']}h`, match rate `{alignment['best_entry_tolerance_0_25']['match_rate']:.4f}`.",
        "- Interpretation: report timestamps and CSV timestamps share the same clock for analysis alignment. If the report is broker GMT+3, then the CSV timestamps should also be treated as broker/server time despite the `+00:00` suffix. For true UTC session labels, subtracting 3 hours may still be appropriate, but price/time joining should use offset `0h`.",
        "",
        "Top offsets for all deals, tolerance 0.25:",
        "",
        "| Offset hours applied to report time | Available | Matched | Match rate | Median distance | P90 distance |",
        "|---:|---:|---:|---:|---:|---:|",
        offset_rows,
        "",
        "## Strategy Tags in Entry Comments",
        "",
        "Entry comment pattern:",
        "",
        "```text",
        "QQ[XAUUSD]1234[T<family>/S<strategy>]",
        "```",
        "",
        "Observed `T/S` counts:",
        "",
        "| Strategy tag | Entry count |",
        "|---|---:|",
        strategy_rows,
        "",
        "Observed `S` counts:",
        "",
        "| S id | Entry count |",
        "|---|---:|",
        s_rows,
        "",
        "Important: only 8 strategy IDs appear in this backtest (`S01`, `S03`, `S04`, `S06`, `S08`, `S09`, `S10`, `S12`). The advertised 12 embedded strategies are not all active in this report/preset, or some had no trades in the tested period.",
        "",
        "## Input Settings Extracted",
        "",
        "| Input | Value |",
        "|---|---|",
    ]
    for key in sorted(summary["inputs"]):
        lines.append(f"| `{key}` | `{summary['inputs'][key]}` |")
    lines.extend(
        [
            "",
            "## CSV Inputs",
            "",
        ]
    )
    for csv in summary["csv"]:
        lines.extend(
            [
                f"### `{csv['file']}`",
                "",
                f"- Rows: `{csv['row_count']}`.",
                f"- Columns: `{csv['columns']}`.",
                f"- Time range: `{csv['time_min']}` to `{csv['time_max']}`.",
                f"- Duplicate timestamps: `{csv['duplicate_timestamps']}`.",
                f"- Missing/invalid timestamps: `{csv['missing_timestamps']}`.",
                "",
            ]
        )
    lines.extend(
        [
            "## Phase 1 Implications",
            "",
            "1. Use `T/S` from comments as the primary strategy label for entries.",
            "2. Reconstruct positions/baskets by matching exit deals to prior open positions because exit comments are blank.",
            "3. Use report timestamp floored to minute for M1 joins; do not apply a `-3h` shift when joining to the provided CSV files.",
            "4. Maintain both `broker_time` and derived `utc_time_est = broker_time - 3h` columns for session analysis until the broker timezone/DST policy is fully confirmed.",
            "5. Because closes/add-ons can happen at second offsets within a minute, M1 management analysis must floor timestamps to minute while preserving original seconds for ordering.",
            "",
        ]
    )
    return "\n".join(lines)


def main() -> None:
    OUTPUTS.mkdir(exist_ok=True)
    workbook = read_workbook_shape()
    raw = pd.read_excel(REPORT_XLSX, sheet_name="Sheet1", header=None)
    order_label_idx = first_row_index(raw, "订单")
    deals_label_idx = first_row_index(raw, "成交")
    metadata, inputs = parse_metadata(raw)
    orders = parse_orders(raw, order_label_idx, deals_label_idx)
    deals = parse_deals(raw, deals_label_idx)
    entries = deals[deals["entry"].eq("in")].copy()
    exits = deals[deals["entry"].eq("out")].copy()
    summary = {
        "workbook": workbook,
        "metadata": {str(k): clean_value(v) for k, v in metadata.items()},
        "inputs": inputs,
        "sections": {
            "orders_label_excel_row": order_label_idx + 1,
            "orders_header_excel_row": order_label_idx + 2,
            "deals_label_excel_row": deals_label_idx + 1,
            "deals_header_excel_row": deals_label_idx + 2,
        },
        "orders": {
            "row_count": int(len(orders)),
            "columns": list(orders.columns),
            "time_min": str(orders["open_time"].min()),
            "time_max": str(orders["open_time"].max()),
            "status_counts": value_counts_dict(orders["status"]),
            "comment_non_null": int(orders["comment"].notna().sum()),
        },
        "deals": {
            "row_count": int(len(deals)),
            "columns": list(deals.columns),
            "time_min": str(deals["time"].min()),
            "time_max": str(deals["time"].max()),
            "type_counts": value_counts_dict(deals["type"]),
            "entry_counts": value_counts_dict(deals["entry"]),
            "entry_count": int(len(entries)),
            "exit_count": int(len(exits)),
            "entries_with_strategy": int(entries["strategy"].notna().sum()),
            "entry_strategy_counts": {str(k): int(v) for k, v in entries["strategy"].value_counts().sort_index().items()},
            "entry_s_counts": {str(int(k)): int(v) for k, v in entries["S"].dropna().astype(int).value_counts().sort_index().items()},
            "entry_t_counts": {str(int(k)): int(v) for k, v in entries["T"].dropna().astype(int).value_counts().sort_index().items()},
        },
        "csv": [csv_summary(M15_CSV), csv_summary(M1_CSV)],
        "time_alignment": alignment_scores(deals),
    }
    (OUTPUTS / "raw_schema_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    (OUTPUTS / "data_audit_report.md").write_text(build_report(summary), encoding="utf-8")
    print("Wrote outputs/data_audit_report.md")
    print("Wrote outputs/raw_schema_summary.json")


if __name__ == "__main__":
    main()
